"""Парсеры файлов: TXT, CSV, XLSX."""

import csv
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from .config import XLSX_MAX_ROWS, XLSX_MAX_COLS, get_logger
from .url_utils import is_valid_url

logger = get_logger(__name__)


def parse_txt(path: str) -> List[str]:
    urls: List[str] = []
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            u = line.strip()
            if is_valid_url(u):
                urls.append(u)
    return urls


def _detect_csv_encoding(path: str) -> str:
    """Кодировка CSV по BOM или первым 4 КБ."""
    with open(path, "rb") as fb:
        head = fb.read(4096)
    if head.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    try:
        head.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        return "cp1251"


def parse_csv(path: str) -> List[str]:
    """Потоковое чтение CSV без полной загрузки в память."""
    urls: List[str] = []
    encoding = _detect_csv_encoding(path)

    with open(path, "r", encoding=encoding, newline="") as f:
        sample = f.readline()
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel

        for row in csv.reader(f, dialect):
            for cell in row:
                c = cell.strip().strip('"').strip("'")
                if is_valid_url(c):
                    urls.append(c)

    return list(dict.fromkeys(urls))


def parse_xlsx(path: str) -> List[str]:
    """Обход XLSX с ограничением по строкам и колонкам."""
    urls: List[str] = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except InvalidFileException as e:
        raise ValueError(f"Повреждённый или несовместимый файл Excel: {e}") from e
    except PermissionError as e:
        raise ValueError("Нет доступа к файлу (возможно, он открыт в Excel)") from e
    except Exception as e:
        raise ValueError(f"Не удалось открыть Excel: {e}") from e

    try:
        for ws in wb.worksheets:
            max_row = min(ws.max_row or 0, XLSX_MAX_ROWS)
            max_col = min(ws.max_column or 0, XLSX_MAX_COLS)
            if max_row == 0 or max_col == 0:
                continue
            for row in ws.iter_rows(
                min_row=1, max_row=max_row,
                min_col=1, max_col=max_col,
                values_only=True,
            ):
                for cell in row:
                    if isinstance(cell, str):
                        c = cell.strip()
                        if is_valid_url(c):
                            urls.append(c)
    finally:
        try:
            wb.close()
        except Exception:
            logger.exception("Ошибка при закрытии xlsx-файла %s", path)
    return list(dict.fromkeys(urls))


def parse_any_file(path: str) -> List[str]:
    """Универсальный парсер по расширению файла."""
    ext = Path(path).suffix.lower()
    if ext == ".txt":
        return parse_txt(path)
    if ext == ".csv":
        return parse_csv(path)
    if ext in (".xlsx", ".xlsm"):
        return parse_xlsx(path)
    raise ValueError(f"Неподдерживаемый формат: {ext}")


__all__ = ["parse_txt", "parse_csv", "parse_xlsx", "parse_any_file"]