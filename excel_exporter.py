"""Генерация Excel-отчёта с двумя листами: результаты и статистика."""

from datetime import datetime
from typing import Callable, List, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import CheckResult


class ExcelExporter:
    """Стили создаются один раз в __init__ и переиспользуются во всём отчёте."""

    def __init__(self) -> None:
        self.header_font = Font(bold=True, color="FFFFFF", size=11)
        self.header_fill = PatternFill("solid", fgColor="2F5597")
        self.center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin = Side(style="thin", color="BFBFBF")
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.fill_ok = PatternFill("solid", fgColor="C6EFCE")
        self.fill_no = PatternFill("solid", fgColor="FFC7CE")
        self.fill_err = PatternFill("solid", fgColor="FFEB9C")
        self.font_ok = Font(color="006100")
        self.font_no = Font(color="9C0006")
        self.font_err = Font(color="9C6500")
        self.link_font = Font(color="0563C1", underline="single")
        self.bold_font = Font(bold=True)

    def write(
        self,
        path: str,
        results: List[CheckResult],
        status_fn: Callable[[Optional[bool]], str],
    ) -> None:
        """
        Создаёт .xlsx-файл. Исключения при сохранении обрабатываются точечно,
        чтобы UI показал пользователю понятное сообщение.
        """
        wb = openpyxl.Workbook()
        self._write_results_sheet(wb.active, results, status_fn)
        self._write_stats_sheet(wb.create_sheet("Статистика"), results)

        try:
            wb.save(path)
        except PermissionError as error:
            raise PermissionError(
                f"Файл «{path}» недоступен для записи — "
                "возможно, он открыт в Excel или нет прав доступа."
            ) from error
        except FileNotFoundError as error:
            raise FileNotFoundError(
                f"Папка для сохранения исчезла: «{path}»"
            ) from error
        except OSError as error:
            raise OSError(
                f"Ошибка файловой системы при сохранении «{path}»: {error}"
            ) from error

    def _write_results_sheet(
        self,
        ws: Worksheet,
        results: List[CheckResult],
        status_fn: Callable[[Optional[bool]], str],
    ) -> None:
        ws.title = "Результаты"
        headers = ["№", "URL", "Yandex", "Google", "Ошибки"]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.font, c.fill, c.alignment, c.border = (
                self.header_font, self.header_fill, self.center, self.border,
            )

        for i, r in enumerate(results, 1):
            row_idx = i + 1
            ws.append([
                i, r["url"],
                status_fn(r["yandex"]),
                status_fn(r["google"]),
                "; ".join(r["errors"]),
            ])
            url_cell = ws.cell(row=row_idx, column=2)
            url_cell.hyperlink = r["url"]
            url_cell.font = self.link_font

            ws.cell(row=row_idx, column=1).alignment = self.center
            ws.cell(row=row_idx, column=2).alignment = self.left
            ws.cell(row=row_idx, column=3).alignment = self.center
            ws.cell(row=row_idx, column=4).alignment = self.center
            ws.cell(row=row_idx, column=5).alignment = self.left
            for col_idx in range(1, 6):
                ws.cell(row=row_idx, column=col_idx).border = self.border

            self._paint_status(
                ws.cell(row=row_idx, column=3), r["yandex"], r["errors"], "Yandex",
            )
            self._paint_status(
                ws.cell(row=row_idx, column=4), r["google"], r["errors"], "Google",
            )
            if r["errors"]:
                ec = ws.cell(row=row_idx, column=5)
                ec.fill, ec.font = self.fill_err, self.font_err

        widths = {1: 6, 2: 60, 3: 22, 4: 22, 5: 45}
        for col_idx, w in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:E{len(results) + 1}"

    def _paint_status(
        self, cell, value: Optional[bool], errors: List[str], label: str,
    ) -> None:
        if value is True:
            cell.fill, cell.font = self.fill_ok, self.font_ok
        elif value is False:
            cell.fill, cell.font = self.fill_no, self.font_no
        elif any(label in e for e in errors):
            cell.fill, cell.font = self.fill_err, self.font_err

    def _write_stats_sheet(
        self, ws: Worksheet, results: List[CheckResult],
    ) -> None:
        def count(pred: Callable[[CheckResult], bool]) -> int:
            return sum(1 for r in results if pred(r))

        y_i = count(lambda r: r["yandex"] is True)
        y_n = count(lambda r: r["yandex"] is False)
        y_e = count(lambda r: r["yandex"] is None)
        g_i = count(lambda r: r["google"] is True)
        g_n = count(lambda r: r["google"] is False)
        g_e = count(lambda r: r["google"] is None)
        total_errors = sum(len(r["errors"]) for r in results)

        rows = [
            ["Параметр", "Значение"],
            ["Дата проверки", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Всего URL", len(results)],
            ["", ""],
            ["Yandex — проиндексировано", y_i],
            ["Yandex — не проиндексировано", y_n],
            ["Yandex — ошибок", y_e],
            ["", ""],
            ["Google — проиндексировано", g_i],
            ["Google — не проиндексировано", g_n],
            ["Google — ошибок", g_e],
            ["", ""],
            ["Всего ошибок запросов", total_errors],
        ]
        for row in rows:
            ws.append(row)

        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment, cell.border = (
                self.header_font, self.header_fill, self.center, self.border,
            )
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
            for cell in row:
                cell.border = self.border
                cell.alignment = self.left if cell.column == 1 else self.center
        for r_idx in (5, 9, 13):
            ws.cell(row=r_idx, column=1).font = self.bold_font

        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 22


__all__ = ["ExcelExporter"]